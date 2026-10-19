#!/usr/bin/env python
# coding: utf-8
"""Genera presentacion6_05_10.pptx: lo hecho despues de la reunion del 5/10, para la siguiente.
Los pedidos de esa reunion: (1) la grilla 64x64 de g = -4 con KMC, en crudo; (2) mas puntos de control
y centrados donde la reconstruccion es mas sensible o mas difiere del continuo; (3) g = -8 a 4 de a 1
con el mismo procedimiento y KMC; mas la comparacion procedimiento vs puntos al azar.

Los g negativos (y todos los g) se muestran con KMC; el continuo aparece solo como referencia o
donde se mide el metodo contra una imagen completa (procedimiento vs azar).

Uso (desde ~/Tesis-autoexp):
    ./venv/bin/python presentaciones/presentacion6_05_10/figuras.py --grid <resultados de la grilla>
    ./venv/bin/python presentaciones/presentacion6_05_10/make_presentacion6_05_10.py
"""
import csv
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(REPO, "presentaciones", "explicacion_autoexp"))
sys.path.insert(0, os.path.join(REPO, "presentaciones", "presentacion5_05_10"))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import make_presentacion5_05_10 as P5  # noqa: E402  (estilo de reunion: titulo azul, texto gris)

M = P5.M
FIG = os.path.join(HERE, "fig")
F4 = os.path.join(REPO, "presentaciones", "presentacion4_autoexp", "fig")
KMC = os.path.join(REPO, "autoexp", "kmc")
OUT = os.path.join(HERE, "presentacion6_05_10.pptx")


def vs_random():
    rows = {float(r["g"]): r for r in csv.DictReader(open(os.path.join(KMC, "comparacion_random", "psnr.csv")))}
    return {g: (float(r["psnr_procedimiento"]), float(r["psnr_azar_media"])) for g, r in rows.items()}


def s_portada(prs):
    s = M.new_slide(prs)
    P5._text(s, 0.7, 2.2, 12.0, 2.0, ["KMC en 13 valores de g y la grilla 64×64 de g = −4"], size=36,
             color=P5.TBLUE)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    P5._text(s, 0.7, 4.0, 12.0, 1.5, ["Lo pedido en la reunión del 5/10: imagen en crudo, control en las zonas "
                                      "críticas, g = −8 a 4 con KMC, y elegir vs. azar  ·  Santiago Afonso"],
             size=19, color=P5.TGREY)


def s_pedidos(prs, grid):
    s = M.new_slide(prs)
    M.title(s, "Lo que quedó de la reunión del 5/10", sub="estado de cada pedido")
    P5.tabla(s, 0.6, 1.65, 12.1, ["pedido", "qué se hizo", "estado"], [
        ["Grilla 64×64 de g = −4 con KMC, en crudo",
         "3392 puntos (log ℓ ≥ −3), 1 corrida por punto, en 4 nodos de Serafín",
         "corriendo: %d terminados" % grid["terminados"]],
        ["Más puntos de control, donde más importa",
         "32 por g (eran 10): 24 sobre la zona crítica + 8 al azar",
         "listo; se corre al cerrar cada g"],
        ["g = −8 a 4, de a 1, con KMC",
         "13 loops con el mismo procedimiento (96 puntos por g)",
         "corriendo"],
        ["¿Sirve elegir los puntos?",
         "procedimiento vs. 96 puntos al azar, misma reconstrucción",
         "hecho (en el continuo)"]], col_w=[3.9, 5.6, 2.6], font=13)
    M.notes(s, """
Los cuatro pedidos de la reunión del 5/10 y en qué están. Todo el KMC corre en Serafín (4 nodos de 64
núcleos reservados por 48 h); los loops que deciden dónde medir corren en Mendieta y mandan las tandas solos.
""")


def s_avance(prs):
    s = M.new_slide(prs)
    M.title(s, "g = −8 a 4 con KMC: avance", sub="96 puntos por g, mismo procedimiento que g = −4 y g = +3")
    M.image(s, os.path.join(FIG, "avance_g.png"), 0.5, 1.55, w=8.4, h=5.7)
    M.card(s, 9.15, 1.7, 3.6, 5.4, "Cómo corre", [
        "Un loop por g: grilla de 30 → bisección sobre el borde → relleno.",
        "Cada tanda es un paso dentro de un nodo de Serafín; al terminar, el loop elige la siguiente.",
        "g = −5 a −8 arrancaron el 5/10 con su grilla de 30.",
        "Al llegar a 96 se corren los 32 de control."], size=13)
    M.notes(s, """
Puntos de KMC terminados por g (sin contar los pseudo-puntos que salen de corridas cortadas por tiempo).
Los negativos más avanzados: g = −4 y g = −2. Los nuevos g = −5 a −8 están en su primera tanda (30 puntos
de la grilla). Las tandas de bisección son chicas (3-4 puntos) porque cada una depende de la anterior.
""")


def s_neg(prs):
    s = M.new_slide(prs)
    M.title(s, "g negativos con KMC: g = −1 a −4",
            sub="arriba el mapa con los puntos de KMC de hoy (x roja: corriendo) · abajo el continuo, solo de referencia")
    M.image(s, os.path.join(FIG, "kmc_neg_1a4.png"), 0.4, 1.6, w=12.5, h=5.7)
    M.notes(s, """
Reconstrucción con lo que hay hoy (mismo método: borde + cero abajo + TPS + monotonía, sin DIP todavía).
Todavía en curso: las zonas con escalones (g = −3) son donde faltan puntos, y es justamente donde el loop
está mandando la bisección. En los cuatro el borde del KMC tiene la misma inclinación que el del continuo.
""")
    s = M.new_slide(prs)
    M.title(s, "g = −5 a −8 con KMC: primera tanda",
            sub="la grilla de 30 puntos (x rojas) corre desde el 5/10 · abajo el continuo de cada g")
    M.image(s, os.path.join(FIG, "kmc_neg_5a8.png"), 0.4, 1.6, w=12.5, h=5.7)
    M.notes(s, """
Los g más negativos son los que pidieron agregar. Se verificó que el KMC acepta g = −8 (parametros.dat).
Primero la grilla de 30 (5 filas × 6 columnas); después el loop sigue solo con la bisección.
""")


def s_pos(prs):
    s = M.new_slide(prs)
    M.title(s, "g = 0 a 4 con KMC", sub="arriba el mapa con los puntos de KMC de hoy · abajo el continuo (referencia)")
    M.image(s, os.path.join(FIG, "kmc_pos.png"), 0.4, 1.6, w=12.5, h=5.4)
    M.notes(s, "Mismo procedimiento; g = +3 es el más avanzado de los positivos.")


def s_grid(prs, grid):
    s = M.new_slide(prs)
    M.title(s, "Grilla 64×64 de g = −4 con KMC, en crudo",
            sub="todos los puntos de la grilla, sin planificador ni reconstrucción · en curso")
    M.image(s, os.path.join(FIG, "grid64_parcial.png"), 0.5, 1.55, w=7.2, h=5.75)
    M.card(s, 8.0, 1.7, 4.75, 5.4, "Cómo se corre", [
        "64 × 64 = 4096 puntos; las columnas con log ℓ < −3 no se corren (una corrida tarda días): "
        "quedan %d." % grid["total"],
        "Se empieza por lo lento (ℓ chico): lo que se ve son las primeras columnas.",
        "Terminados: %d · corriendo: %d." % (grid["terminados"], grid["corriendo"]),
        "Para qué: ver el mapa sin interpolar y tener una imagen completa de KMC para medir las "
        "reconstrucciones de 96 puntos."], size=13)
    M.notes(s, """
Pedido de la reunión: correr una imagen entera de g = −4 a 64 × 64 (no 128) para ver el diagrama en crudo.
3392 corridas, 1 réplica por punto, repartidas en 4 nodos de Serafín con prioridad baja para no frenar los
loops. Arriba a la izquierda ya aparece la franja de log Ξ ≥ 1.4 con SoC ≈ 0 (la anomalía de Ξ alto que
consultamos); abajo del borde da 0 como en el continuo.
Cuando termine: mapa crudo, continuo, diferencia; y el error de la reconstrucción de 96 puntos contra esta
imagen (la primera medida de error contra KMC completo, no contra el continuo).
""")


def s_control(prs):
    s = M.new_slide(prs)
    M.title(s, "Control nuevo: 32 puntos, centrados donde más importa",
            sub="ejemplo con los puntos de hoy de g = −4 · se corren con KMC y se comparan con el mapa")
    M.image(s, os.path.join(FIG, "control_critico_g-4.png"), 0.5, 1.55, w=7.4, h=5.75)
    M.card(s, 8.2, 1.7, 4.55, 5.4, "Cómo se eligen", [
        "Puntaje crítico = cuánto difiere la reconstrucción del continuo + cuánto cambia el mapa (gradiente).",
        "24 puntos (x) sorteados con peso en ese puntaje: cubren todo el frente, no solo el máximo.",
        "8 puntos (+) al azar en todo el mapa.",
        "El error se informa por separado: en la zona crítica y al azar."], size=13)
    M.notes(s, """
Pedido: más de 10 puntos de control, y en los lugares más sensibles o que más difieren del continuo.
Ahora son 32 por g: 75 % sorteados con peso (puntaje)^2, a 4 píxeles como mínimo de los medidos y entre sí;
el resto al azar. log ℓ ≥ −2.7 para que terminen. Con 10 al azar en g = −4 el RMS era 0.027; con este
esquema el error en la zona crítica va a ser mayor, que es lo honesto.
""")


def s_vs_random(prs):
    vr = vs_random()
    neg = [vr[g][0] - vr[g][1] for g in vr if g < 0]
    pos = [vr[g][0] - vr[g][1] for g in vr if g >= 2]
    s = M.new_slide(prs)
    M.title(s, "¿Sirve elegir los puntos? Procedimiento vs. el mismo número al azar",
            sub="medido en el continuo, el único lugar con imagen completa · 96 puntos, ruido tipo KMC, misma reconstrucción")
    M.image(s, os.path.join(F4, "vs_random.png"), 0.4, 1.6, w=12.5, h=4.6)
    M.text(s, 0.6, 6.35, 12.1, 0.9, [
        [("Con g < 0 elegir suma %.0f a %.0f dB; con g ≥ 2, %.0f a %.0f dB. " % (min(neg), max(neg), min(pos), max(pos)),
          {"bold": True}),
         ("Con KMC se repite en g = −4 cuando esté la grilla 64×64.", {})]], size=15)
    M.notes(s, """
El método se mide contra el continuo porque ahí está la imagen completa. 10 sorteos al azar por g.
Al azar con TPS común da lo mismo (±0.4 dB): la ganancia viene de dónde caen los puntos, no de la
reconstrucción. Con g negativo el frente es abrupto y al azar casi nadie cae cerca: el azar queda en ~19-20 dB
desde g = −4 para abajo, el procedimiento en ~51.
Con KMC: con la grilla 64×64 de g = −4 se puede hacer lo mismo sorteando 96 puntos de la grilla.
""")


def s_sigue(prs):
    s = M.new_slide(prs)
    M.title(s, "Cómo sigue")
    M.text(s, 0.6, 1.6, 12.1, 5.6, [
        [("Nodos: ", {"bold": True}), ("los 4 de Serafín vencen el 7/10; ya están pedidos los 4 de relevo y "
                                       "los loops y la grilla pasan solos.", {})],
        [("Cada g al llegar a 96: ", {"bold": True}), ("mapa final (con DIP) y control de 32 puntos.", {})],
        [("Grilla 64×64: ", {"bold": True}), ("mapa crudo, diferencia con el continuo y error de la "
                                              "reconstrucción de 96 puntos contra KMC completo.", {})],
        [("Para la reunión: ", {"bold": True}), ("los 13 mapas de KMC lado a lado y cómo cambia el borde con g.", {})]],
        size=17, bullets=True, space=12)
    M.notes(s, "Próximos pasos.")


def main():
    grid = json.load(open(os.path.join(FIG, "grid64_parcial.json")))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(M.SW), Inches(M.SH)
    s_portada(prs)
    s_pedidos(prs, grid)
    s_avance(prs)
    s_neg(prs)
    s_pos(prs)
    s_grid(prs, grid)
    s_control(prs)
    s_vs_random(prs)
    s_sigue(prs)
    prs.save(OUT)
    print("->", OUT, "(%d filminas)" % len(prs.slides))


if __name__ == "__main__":
    main()
