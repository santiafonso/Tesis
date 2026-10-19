#!/usr/bin/env python
# coding: utf-8
"""Genera presentacion5_05_10.pptx: la presentacion corta para la reunion del 5/10. Un solo hilo:
1. Optuna (hiperparametros de DIP) -> 2. interpolacion (Andres) -> 3. estructura fisica: la linea
(Maxi; biseccion + recta + monotonia) -> 4. DIP para suavizar; despues resultados: continuo y las
dos reconstrucciones con KMC real (g = -4 y g = +3) con su control de 10 puntos al azar.

Reusa las figuras de explicacion_autoexp (figuras.py / figuras_kmc.py), presentacion4_autoexp
y autoexp/kmc/g*/ (dip/final.png, vs_continuo.png, control.png). No recalcula nada.

Uso (desde ~/Tesis-autoexp):
    ./venv/bin/python presentaciones/presentacion5_05_10/make_presentacion5_05_10.py
"""
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(REPO, "presentaciones", "explicacion_autoexp"))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import make_explicacion_autoexp as M  # noqa: E402  (helpers de estilo: misma paleta y tipografia)
from figuras_kmc import CAP_PASOS  # noqa: E402

FE = os.path.join(REPO, "presentaciones", "explicacion_autoexp", "fig")
F4 = os.path.join(REPO, "presentaciones", "presentacion4_autoexp", "fig")
KMC = os.path.join(REPO, "autoexp", "kmc")
OUT = os.path.join(HERE, "presentacion5_05_10.pptx")

# ---- estilo de las presentaciones de reunion (1-4): fondo blanco, titulo azul en negrita, texto gris,
# sin tarjetas de color. Se pisan los helpers de explicacion_autoexp; el contenido no cambia.
TBLUE, TGREY = "1F3B73", "404040"
_text, _box = M.text, M.box


def _new_slide(prs, dark=False):
    return _new_slide.orig(prs, dark=False)


_new_slide.orig = M.new_slide


def _text_gris(s, x, y, w, h, paras, size=16, color=None, **kw):
    if color in (None, M.INK, M.MUTED, M.WHITE, "C9D1E0", "8A94A8"):
        color = TGREY
    return _text(s, x, y, w, h, paras, size=size, color=color, **kw)


def _title(s, t, step=None, dark=False, sub=None):
    _text(s, 0.6, 0.35, 12.1, 0.9, [t], size=27, color=TBLUE, space=0)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    if sub:
        _text(s, 0.6, 1.1, 12.1, 0.4, [[(sub, {"italic": True})]], size=13, color=TGREY, space=0)


def _card(s, x, y, w, h, head, body, tint=None, head_color=None, size=14):
    _text(s, x + 0.1, y + 0.1, w - 0.2, h - 0.2,
          [[(head, {"bold": True, "color": TBLUE, "size": size + 2})]] + body, size=size, color=TGREY,
          space=4)


def _box_linea(s, x, y, w, h, fill, radius=True):
    sh = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # rectangulo, sin relleno
    sh.fill.background()
    sh.line.color.rgb = M.rgb("BFBFBF")
    sh.shadow.inherit = False
    return sh


def _stat(s, x, y, w, big, label, color=None, big_size=40, dark=False):
    _text(s, x, y, w, 0.8, [big], size=big_size, color=TBLUE, space=0)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    _text(s, x, y + 0.82, w, 0.7, [label], size=13, color=TGREY, space=0)


def _arrow(s, x, y, w=0.55, h=0.4, color=None):
    return _arrow.orig(s, x, y, w, h, color="A6A6A6")


_arrow.orig = M.arrow
M.new_slide, M.text, M.title, M.card, M.box, M.stat, M.arrow = (_new_slide, _text_gris, _title, _card,
                                                                _box_linea, _stat, _arrow)


def ctrl(g):
    return json.load(open(os.path.join(KMC, g, "control.json")))


def s_portada(prs):
    s = M.new_slide(prs)
    _text(s, 0.7, 2.2, 12.0, 2.0, ["Reconstruir el diagrama de fases con 64 puntos de KMC"], size=36,
          color=TBLUE)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    _text(s, 0.7, 4.0, 12.0, 1.5, ["TPS + bisección + borde + monotonía + DIP · primeros resultados con KMC real "
                                   "(g = −4 y g = +3)  ·  Santiago Afonso  ·  5/10/2026"], size=19, color=TGREY)


def tabla(s, x, y, w, headers, rows, col_w, font=13):
    tbl = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w),
                             Inches(0.42 * (len(rows) + 1))).table
    for c, cw in enumerate(col_w):
        tbl.columns[c].width = Inches(cw)
    for r, row in enumerate([headers] + rows):
        for c, v in enumerate(row):
            p = tbl.cell(r, c).text_frame.paragraphs[0]
            p.text = str(v)
            p.font.size = Pt(font + (1 if r == 0 else 0))
            p.font.bold = r in (0, len(rows))  # encabezado y la mejor de Optuna


def s_optuna(prs):
    s = M.new_slide(prs)
    M.title(s, "1. Optuna: los hiperparámetros de DIP",
            sub="pedido de la reunión del 17/9 · 181 pruebas en Mendieta · 64 puntos, 3 valores de g")
    tabla(s, 0.6, 1.65, 12.1, ["configuración (64 puntos en grilla)", "PSNR medio [dB]", "peor g [dB]"],
          [["interpolación TPS, como referencia", "24.4", "21.2"],
           ["DIP, receta de la tesis (LR 1e-3, reg 0.08, skip 4)", "25.0", "21.9"],
           ["DIP, configuración 'vase' del paper", "28.5", "25.3"],
           ["DIP, la mejor de Optuna (181 pruebas)", "29.9", "26.2"]], col_w=[7.7, 2.2, 2.2])
    M.text(s, 0.6, 4.1, 12.1, 3.0, [
        [("Qué buscó: ", {"bold": True}), ("tasa de aprendizaje, ruido de regularización, arquitectura, "
                                           "iteraciones, tipo de entrada, y el muestreo (grilla o uniforme, "
                                           "densidad y corrimiento de la grilla) como un hiperparámetro más.", {})],
        [("Resultado: ", {"bold": True}), ("la mejor configuración mejora ~5 dB, pero con 64 puntos DIP no pasa "
                                           "los 30 dB. La grilla siempre le gana al muestreo uniforme.", {})],
        [("Conclusión: ", {"bold": True}), ("ajustar DIP no alcanza; hay que cambiar dónde se mide y usar lo que "
                                            "sabemos del mapa. Eso es lo que sigue.", {})]],
        size=15, bullets=True, space=8)
    M.notes(s, """
Pedido 1-3 de la reunión del 17/9: Optuna (búsqueda automática de hiperparámetros), grid/uniform tocando
hiperparámetros, y la grilla como hiperparámetro. Estudio profes_grid_uniform: 181 trials en Mendieta.
Los mejores son todos la config 'vase' del paper afinada: entrada meshgrid, skip 0, LR 0.005-0.017,
~3000 iteraciones. Uniform nunca aparece arriba. Mejor: 29.9 dB de media (peor g 26.2) con 64 puntos.
Un segundo estudio (28 trials) con el método final como objetivo tampoco mejoró la receta de la tesis:
la configuración de DIP ya está cerca del óptimo. No hubo una mejora sustancial.
""")


def s_continuo96(prs):
    s = M.new_slide(prs)
    M.title(s, "Con 96 puntos, sin DIP: los 9 g sobre 40 dB",
            sub="arriba la imagen real, al medio la reconstrucción, abajo el error · 9 valores de g")
    M.image(s, os.path.join(F4, "mejores_9g_96pts.png"), 0.4, 1.65, w=12.5, h=5.6)
    M.notes(s, """
Misma receta (grilla 30 + bisección + relleno, línea + 0 abajo + TPS + monotonía), sin DIP, 96 puntos
(0.6 % de la imagen): 44.4 dB de media, peor g 40.2 (g = -2). Corrida VAL9_96pts.
Por eso el KMC pasa a 96 puntos en todos los g: el costo extra es chico y el margen grande.
""")


def s_problema(prs):
    s = M.new_slide(prs)
    M.title(s, "El problema: cada punto de KMC cuesta horas")
    M.image(s, os.path.join(FE, "problema.png"), 0.6, 1.4, w=7.6, h=5.6)
    M.card(s, 8.5, 1.6, 4.25, 5.2, "La idea", [
        "Correr KMC en pocos puntos (64) del plano (log Ξ, log ℓ).",
        "Reconstruir el mapa completo de 128 × 128 a partir de esos puntos.",
        "Primero se afinó todo con el modelo del continuo, donde sí se conoce la imagen completa y se "
        "puede medir el error."], size=15)
    M.notes(s, """
Una corrida de KMC tarda de minutos a días según log ℓ. Un mapa de 128 × 128 son 16 mil corridas: imposible.
Entonces: medir pocos puntos bien elegidos y reconstruir el resto. Todo el método se armó y se midió
primero contra el modelo del continuo (galpynostatic), donde tenemos la imagen completa para comparar.
""")


def s_scipy(prs):
    s = M.new_slide(prs)
    M.title(s, "2. Interpolar, como sugirió Andrés: la TPS",
            sub="mismos 64 puntos en grilla · arriba la reconstrucción, abajo el error")
    M.image(s, os.path.join(FE, "scipy_metodos.png"), 0.4, 1.7, w=12.5, h=4.6)
    M.text(s, 0.6, 6.45, 12.1, 0.8, [
        [("TPS (thin-plate spline): ", {"bold": True}),
         ("la superficie más lisa que pasa por los puntos. Pero sola no alcanza: todas fallan en el borde, "
          "donde el SoC cae a 0 de golpe.", {})]], size=15)
    M.notes(s, """
Lo que sugirió Andrés en la reunión del 17/9: interpolar con scipy. Probamos los cuatro métodos con los mismos
64 puntos: vecino más cercano, lineal, cúbica y TPS (scipy.interpolate.RBFInterpolator, kernel
thin_plate_spline). Media en 3 g: 19.1 / 21.9 / 22.9 / 24.4 dB. La TPS es la mejor y es la base.
Pero mirá la fila de abajo: todo el error está en el borde inclinado. Ninguna interpolación puede hacer un
salto seco con puntos a 16 píxeles de distancia. De ahí salen las dos ideas siguientes.
""")


def s_biseccion(prs):
    s = M.new_slide(prs)
    M.title(s, "3. La estructura física (idea de Maxi): la línea",
            sub="dónde medir: grilla y después bisección sobre el borde")
    M.image(s, os.path.join(FE, "muestreo_etapas.png"), 0.4, 1.4, w=12.5, h=4.2)
    M.text(s, 0.6, 5.8, 12.1, 1.4, [
        [("Grilla (30 puntos): ", {"bold": True}), ("ver la forma general.", {})],
        [("Bisección: ", {"bold": True}), ("en cada columna donde el SoC pasa de alto a 0, búsqueda binaria "
                                           "hasta ubicar el salto con precisión de 1 píxel.", {})],
        [("Relleno: ", {"bold": True}), ("los últimos puntos van donde el modelo más se equivoca.", {})]],
        size=15, bullets=True, space=4)
    M.notes(s, """
La información está en el borde, no en las mesetas. Entonces: grilla gruesa de 30 puntos para ver la forma;
en cada columna de la grilla donde el SoC salta de alto a 0, se mide en el medio, después en el medio de la
mitad que contiene el salto, etc. (búsqueda binaria). Con ~5 rondas el borde queda ubicado a 1 píxel.
Lo que sobra (relleno) va donde, sacando cada punto y prediciéndolo con los demás, el error es mayor.
Con KMC esto se corre por tandas: cada ronda de bisección es una tanda de corridas en paralelo.
""")


def s_recta(prs):
    s = M.new_slide(prs)
    M.title(s, "3. La estructura física: la línea y la monotonía",
            sub="mismos 64 puntos en las cuatro columnas (continuo, g = −2)")
    M.image(s, os.path.join(FE, "recon_pasos.png"), 0.5, 1.65, w=8.2, h=5.6)
    M.card(s, 9.0, 1.7, 3.75, 2.75, "La recta que corta", [
        "Con los pares vecinos 0 / no-0 de la bisección se traza el borde (ajuste robusto: un punto mal "
        "medido no la mueve).",
        "Debajo: 0. Arriba: TPS solo con los puntos de arriba."], size=12.5)
    M.card(s, 9.0, 4.6, 3.75, 2.6, "Monotonía", [
        "Físicamente el SoC no sube al aumentar ℓ ni Ξ.",
        "Se fuerza eso en el mapa: borra las ondulaciones de la TPS."], tint=M.GREEN_T, head_color="137A55",
           size=12.5)
    M.notes(s, """
De izquierda a derecha, misma entrada: TPS común; + cero debajo del borde; + rampa alineada al borde;
+ monotonía. Cada idea suma dB (ver los títulos).
La recta: los puntos de la bisección dan pares vecinos donde uno vale 0 y el otro no; el borde pasa por el
medio. Por esos ~5 puntos se ajusta una recta con RANSAC: se prueban rectas por pares de puntos y gana la
que deja más puntos cerca, así un punto mal detectado no la tuerce (si preguntan: "ajuste robusto").
Debajo de la recta el mapa vale 0 exacto; arriba se interpola solo con los puntos de arriba (si la TPS
usara también los ceros, los "chorrearía" hacia la meseta).
Monotonía: el SoC no sube ni al bajar en el mapa ni al ir a la derecha. Cada punto medido acota una esquina
del mapa, y al final se proyecta al mapa monótono más cercano. Es conocimiento físico, no de la imagen.
""")


def s_dip(prs):
    s = M.new_slide(prs)
    M.title(s, "4. DIP para suavizar: se mezcla con la TPS por zonas")
    M.image(s, os.path.join(F4, "zonas_tps_dip.png"), 0.6, 1.4, w=12.1, h=3.7)
    M.card(s, 0.6, 5.3, 6.0, 1.85, "Cerca del borde manda la TPS", [
        "Es precisa en el salto, donde DIP deja escalones."], size=14)
    M.card(s, 6.75, 5.3, 6.0, 1.85, "Lejos del borde manda DIP", [
        "Es más lisa en las mesetas, donde la TPS ondula. Si DIP no reproduce los puntos medidos, se descarta."],
           tint=M.ORANGE_T, head_color="B4471C", size=14)
    M.notes(s, """
Cada uno es bueno en una zona. Se mezclan con un peso que depende de la distancia al borde (20 px):
pegado al borde TPS, lejos DIP. A DIP se le dan los puntos medidos + ceros debajo del borde.
Guarda: si DIP no ajusta los puntos medidos (error > 0.02 + 2σ), se usa solo la TPS.
En el continuo, 64 puntos, 9 g: TPS sola ~40.6 dB, TPS + DIP ~41.3 dB; con ruido DIP suma más.
""")


def s_metodo(prs):
    s = M.new_slide(prs)
    M.title(s, "Todo junto: la receta", sub="los 8 pasos con KMC real están en el apéndice")
    pasos = [("Medir", "64 puntos de KMC: grilla + bisección sobre el borde", M.BLUE_T, M.NAVY),
             ("Interpolar", "TPS: la superficie más lisa por los puntos (la mejor de scipy)", M.BLUE_T, M.NAVY),
             ("Física", "debajo del borde SoC = 0, y el SoC no sube con ℓ ni con Ξ", M.GREEN_T, "137A55"),
             ("DIP", "alisa las mesetas; cerca del borde manda la TPS", M.ORANGE_T, "B4471C")]
    for k, (h, b, t, c) in enumerate(pasos):
        x = 0.6 + k * 3.1
        M.box(s, x, 1.8, 2.75, 1.9, t)
        M.text(s, x + 0.2, 1.9, 2.4, 1.8, [[(h, {"bold": True, "size": 20, "color": TBLUE})], b], size=14, space=4)
        if k < 3:
            M.arrow(s, x + 2.78, 2.55, w=0.3, h=0.35)
    M.image(s, os.path.join(FE, "kmc_pasos_g-4.png"), 2.2, 3.95, w=8.9, h=3.35)
    M.notes(s, """
En una frase: medimos pocos puntos donde está la información (el borde), interpolamos con TPS respetando lo
que sabemos de la física (cero debajo del borde, monotonía) y DIP alisa el resto.
La figura de abajo es el recorrido completo con KMC real (g = −4): puntos, TPS, borde, cero abajo,
monotonía, DIP, peso de la mezcla, final. Si preguntan por un paso, está en el apéndice.
""")


def s_apendice(prs):
    s = M.new_slide(prs)
    _text(s, 0.7, 2.6, 12.0, 1.2, ["Apéndice: cómo funciona cada paso"], size=36, color=TBLUE)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    _text(s, 0.7, 3.9, 12.0, 1.0, ["el problema · los 8 pasos de la reconstrucción con KMC real"], size=19, color=TGREY)


def s_continuo(prs):
    s = M.new_slide(prs)
    M.title(s, "Primero, en el modelo del continuo: 41 dB con 64 puntos",
            sub="arriba la imagen real, al medio la reconstrucción, abajo el error · 9 valores de g")
    M.image(s, os.path.join(F4, "mejores_9g.png"), 0.4, 1.65, w=12.5, h=5.6)
    M.notes(s, """
Validación en los 9 g negativos del continuo con 64 puntos: 41.3 dB de media (peor g 37.7).
Referencias: interpolación de scipy sola ~24 dB; DIP solo con 64 puntos en grilla ~30 dB.
Con 84 puntos los 9 g pasan los 40 dB.
""")


def s_pasos(prs):
    for k, cap in enumerate(CAP_PASOS, 1):
        s = M.new_slide(prs)
        M.title(s, "Con KMC real, g = −4: de los 64 puntos al mapa final", sub="paso %d de 8" % k)
        M.image(s, os.path.join(FE, "kmc_pasos_g-4_%d.png" % k), 0.5, 1.6, w=12.3, h=5.0)
        M.text(s, 0.6, 6.7, 12.1, 0.6, [cap], size=14, color=M.INK, space=0)
        M.notes(s, cap + "\n\nLa misma receta de las filminas anteriores, ahora con puntos de KMC real.")
        M.fade(s)


def s_resultado(prs, g, titulo, nota):
    c = ctrl(g)
    s = M.new_slide(prs)
    M.title(s, titulo, sub="KMC reconstruido · continuo con el mismo g · diferencia")
    M.image(s, os.path.join(KMC, g, "vs_continuo.png"), 0.4, 1.65, w=12.5, h=4.1)
    M.text(s, 0.6, 5.95, 12.1, 1.3, nota, size=15, bullets=True, space=4)
    M.notes(s, "Control: RMS %.3f (TPS + DIP), %.3f (TPS sola).\n" % (c["fusion"]["rms"], c["tps"]["rms"])
            + "\n".join(x if isinstance(x, str) else "".join(t for t, _ in x) for x in nota))
    s = M.new_slide(prs)
    M.title(s, "Control (g = %s): 10 puntos al azar que no se usaron" % {"g-4": "−4", "g3": "+3"}[g],
            sub="se corren con KMC después y se comparan con lo que predice el mapa")
    M.image(s, os.path.join(KMC, g, "control.png"), 0.4, 1.65, w=8.6, h=5.6)
    M.stat(s, 9.3, 1.9, 3.5, "%.3f" % c["fusion"]["rms"], "error RMS de SoC\n(TPS + DIP)", color=M.GREEN)
    M.stat(s, 9.3, 3.55, 3.5, "%.3f" % c["fusion"]["max"], "error máximo", color=M.NAVY)
    M.stat(s, 9.3, 5.2, 3.5, "%.3f" % c["tps"]["rms"], "error RMS con\nla TPS sola", color=M.MUTED)
    M.notes(s, """
Después de reconstruir, se corrieron 10 puntos nuevos al azar (que el método no vio). Lo predicho por el mapa
y lo medido por KMC casi coinciden: los puntos caen sobre la diagonal.
Error RMS %.3f de SoC con TPS + DIP (máximo %.3f); %.3f con la TPS sola. Para comparar: usar el mapa del
continuo como si fuera el KMC da ~0.14 en g = −4.
Limitación honesta: los puntos de control están en log ℓ > −2.2 (más a la izquierda las corridas no terminan
en 8 h). Hay un control extra corriendo en la zona izquierda de g = +3.
""" % (c["fusion"]["rms"], c["fusion"]["max"], c["tps"]["rms"]))


def s_kmc_como(prs):
    s = M.new_slide(prs)
    M.title(s, "Cómo estamos corriendo los puntos de KMC", sub="para revisar con ustedes: ¿está bien así?")
    M.text(s, 0.6, 1.6, 7.3, 5.6, [
        [("Programa: ", {"bold": True}), ("KMC galvanostático (C++, OpenMP), red CS-40×20×40.", {})],
        [("Cada punto: ", {"bold": True}), ("./kmc_fast  Ξ  ℓ  réplica  g  →  descarga hasta el potencial de "
                                            "corte; el valor es el SoC final.", {})],
        [("Réplicas: ", {"bold": True}), ("varias por punto (los núcleos del nodo / cantidad de puntos), con "
                                          "semillas distintas; se promedian. Ruido entre réplicas: σ ≈ 0.005 "
                                          "(0.02 en g = −4).", {})],
        [("Dónde: ", {"bold": True}), ("Serafín, nodo entero de 64 núcleos, 1 núcleo por corrida; las tandas "
                                       "salen solas del planificador (Mendieta coordina).", {})],
        [("Tope: ", {"bold": True}), ("8 h por corrida. Si no termina, el último SoC es solo un mínimo "
                                      "(\"vale al menos\").", {})],
        [("Si no hay filas: ", {"bold": True}), ("el potencial pasa el corte en el primer intervalo → SoC = 0.", {})]],
        size=15, bullets=True, space=8)
    M.card(s, 8.2, 1.65, 4.55, 5.5, "Dudas", [
        "¿El tamaño de red y el corte de potencial son los adecuados?",
        "¿Cuántas réplicas hacen falta por punto?",
        "log ℓ < −3.3: no termina ni en 48 h. ¿Se puede acelerar, o no hace falta esa zona?",
        "log Ξ ≥ 1.4: da menos SoC que justo debajo (en log ℓ = −3.5, SoC ≈ 0 en minutos). ¿Física o algo del código?"],
           tint=M.ORANGE_T, head_color="B4471C", size=13.5)
    M.notes(s, """
Cómo se corre cada punto: el programa de KMC galvanostático (KMC-Galvanostatic_noclus_param_claude.cpp,
compilado como kmc_fast) recibe Ξ, ℓ, el número de réplica y g. Usa la red CS-40x20x40.xyz. Descarga
hasta el potencial de corte y escribe SoC, E y tiempo por intervalo; el valor del punto es el SoC final.
Cada tanda (los puntos que pide el planificador) corre en un nodo de Serafín: 1 núcleo por corrida y
réplicas para llenar el nodo (64 núcleos / puntos). Las réplicas difieren en la semilla (srand(time),
por eso se lanzan con 1 s de diferencia) y se promedian; su dispersión da el ruido σ.
Tope de 8 h: si una corrida no termina, su último SoC es un mínimo (cota). Esas no se usan como valor.
Preguntas para ellos: tamaño de red, corte, número de réplicas, la zona de ℓ chico (no termina ni en
48 h) y la anomalía de Ξ alto.
""")


def s_preguntas(prs):
    s = M.new_slide(prs)
    M.title(s, "Lo que queremos consultar")
    qa = [
        ("¿Es esperable?", "g = −4: el borde del KMC está ~0.4 décadas de Ξ más abajo que en el continuo. "
                          "g = +3: la transición del KMC es más ancha y empieza antes."),
        ("Ξ alto", "En la franja log Ξ ≥ 1.4 el KMC da un SoC menor que justo debajo, en todos los g. "
                   "¿Física o algo de la simulación?"),
        ("ℓ chico", "Con log ℓ < −3.3 las corridas no terminan ni en 48 h. ¿Necesitamos esa zona? "
                    "¿Se pueden acelerar?"),
        ("Siguiente", "Los otros 7 g corren solos en el cluster. Los nodos reservados de Serafín vencen hoy "
                      "a la tarde."),
    ]
    for k, (q, a) in enumerate(qa):
        x = 0.6 + (k % 2) * 6.2
        y = 1.55 + (k // 2) * 2.85
        M.card(s, x, y, 5.95, 2.6, q, [a], tint=[M.BLUE_T, M.ORANGE_T, M.GREEN_T, M.BLUE_T][k],
               head_color=[M.NAVY, "B4471C", "137A55", M.NAVY][k], size=15)
    M.notes(s, "Preguntas para los directores.")


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(M.SW), Inches(M.SH)
    s_portada(prs)
    s_optuna(prs)        # 1. Optuna
    s_scipy(prs)         # 2. interpolacion (Andres)
    s_biseccion(prs)     # 3. estructura fisica (Maxi): donde medir
    s_recta(prs)         #    trazar la linea + monotonia
    s_dip(prs)           # 4. DIP para suavizar
    s_metodo(prs)        # resumen y resultados
    s_continuo(prs)
    s_continuo96(prs)
    s_resultado(prs, "g-4", "KMC real, g = −4", [
        "Misma forma que el continuo (meseta, rampa, salto a 0).",
        "El borde del KMC está más abajo: franja roja (el KMC tiene SoC donde el continuo ya da 0)."])
    s_resultado(prs, "g3", "KMC real, g = +3", [
        "Llega a 0 en el mismo lugar que el continuo.",
        "Pero se descarga antes y más gradual: franja azul (la rampa empieza ~1 década de Ξ antes)."])
    s_kmc_como(prs)
    s_preguntas(prs)
    s_apendice(prs)
    s_problema(prs)
    s_pasos(prs)
    prs.save(OUT)
    print("->", OUT, "(%d filminas)" % len(prs.slides))


if __name__ == "__main__":
    main()
