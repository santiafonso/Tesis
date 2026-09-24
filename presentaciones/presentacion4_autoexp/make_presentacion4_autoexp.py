#!/usr/bin/env python
# coding: utf-8
"""Genera presentacion4_autoexp.pptx (sin fecha, se va actualizando): el loop de
experimentacion automatica (branch `autoexp`) para reconstruir el diagrama de fases con
64 puntos, paso por paso, con graficos. Tambien genera `resumen.png` (una imagen con el
estado actual).

Lee los puntajes de autoexp/runs/<corrida>/score.json (los escribe autoexp.eval), asi que
para actualizarla alcanza con volver a correrlo. Las tablas de DIP en el cluster y los
descartes vienen de autoexp/experimentos.md (bitacora) y estan escritas aca.

Uso (desde ~/Tesis-autoexp):
    ./venv/bin/python presentaciones/presentacion4_autoexp/make_presentacion4_autoexp.py
"""
import json
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT_DIR = os.path.join(REPO, "presentaciones", "presentacion4_autoexp")
PRES_NAME = "presentacion4_autoexp"
RUNS = os.path.join(REPO, "autoexp", "runs")
FIG = os.path.join(OUT_DIR, "fig")
import sys

sys.path.insert(0, REPO)
from autoexp.oracle import load_truth  # noqa: E402

DEV = ["-4.0", "-2.0", "-0.5"]
GS9 = ["-4.0", "-3.5", "-3.0", "-2.5", "-2.0", "-1.5", "-1.0", "-0.5", "0.0"]

# paleta categorica validada (dataviz, references/palette.md), en orden fijo
C1, C2, C3 = "#2a78d6", "#eb6834", "#1baf7a"
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
SLIDE_W, SLIDE_H = 13.33, 7.5
BLUE = RGBColor(0x1F, 0x3B, 0x73)
GREY = RGBColor(0x40, 0x40, 0x40)

plt.rcParams.update({
    "font.size": 11, "axes.edgecolor": INK2, "axes.labelcolor": INK, "xtick.color": INK2,
    "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8, "axes.axisbelow": True,
    "figure.facecolor": "white", "axes.facecolor": "white", "legend.frameon": False,
})


# ------------------------------------------------------------------ datos -------
def score(run):
    p = os.path.join(RUNS, run, "score.json")
    return json.load(open(p)) if os.path.exists(p) else None


def per_g(run):
    s = score(run)
    return {str(r["g"]): r["psnr"] for r in s["per_g"]} if s else {}


def dev_stats(run):
    d = per_g(run)
    v = [d[g] for g in DEV if g in d]
    return (float(np.mean(v)), float(np.min(v))) if len(v) == len(DEV) else (np.nan, np.nan)


# Recorrido (siempre medido en los 3 g de desarrollo, mismo evaluador, 64 puntos)
STEPS = [
    ("Grilla 8×8\n+ TPS (scipy)", "base_grid8x8_rbf_tps"),
    ("Muestreo\nadaptativo", "adapt_n148_p1_tps"),
    ("Bisección\ndel frente", "bisect36_front_split"),
    ("Acantilado\na cero", "cliff3_b36_cm0.03"),
    ("Rampa alineada\nal borde", "cliffA_al0.5_b15"),
    ("Relleno por\ngradiente", "cliffG_n36_cliff_gap4"),
    ("TPS con y\ncomprimida", "cliffV_0.6"),
    ("Relleno por\nvalidación cruzada", "cliffL_n30_b4"),
    ("Monotonía", "cliffM_it100"),
    ("Franja adaptativa\n+ cotas", "bounds_true"),
    ("Fusión\nTPS + DIP", "fuse2_hib_mid_B10"),
]


# --------------------------------------------------------------- figuras -------
def _save(fig, name):
    os.makedirs(FIG, exist_ok=True)
    p = os.path.join(FIG, name)
    fig.savefig(p, dpi=130, bbox_inches="tight")
    plt.close(fig)
    return p


def fig_progress(ax=None):
    labels = [l for l, _ in STEPS]
    st = [dev_stats(r) for _, r in STEPS]
    mean = np.array([s[0] for s in st])
    mn = np.array([s[1] for s in st])
    own = ax is None
    if own:
        fig, ax = plt.subplots(figsize=(13, 5.2))
    x = np.arange(len(labels))
    w = 0.38
    ax.bar(x - w / 2 - 0.01, mean, w, color=C1, label="PSNR medio (3 g)")
    ax.bar(x + w / 2 + 0.01, mn, w, color=C2, label="PSNR del peor g")
    for i in range(len(x)):
        ax.text(x[i] - w / 2, mean[i] + 0.4, "%.1f" % mean[i], ha="center", fontsize=8.5, color=INK)
        ax.text(x[i] + w / 2, mn[i] + 0.4, "%.1f" % mn[i], ha="center", fontsize=8.5, color=INK2)
    ax.axhline(38, color=INK, ls=(0, (4, 3)), lw=1)
    ax.text(-0.45, 38.4, "objetivo 38 dB", ha="left", fontsize=9, color=INK)
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=8.5)
    ax.set_ylim(15, 47)
    ax.set_ylabel("PSNR [dB]")
    ax.grid(axis="x", visible=False)
    ax.legend(loc="upper left", ncol=2, bbox_to_anchor=(0.0, 1.0))
    ax.set_title("Cada paso del recorrido, con los mismos 64 puntos de presupuesto", loc="left", fontsize=12)
    if own:
        return _save(fig, "progreso.png")


def fig_baselines():
    runs = [("vecino más\ncercano", "base_grid8x8_nearest"), ("lineal", "base_grid8x8_linear"),
            ("cúbica", "base_grid8x8_cubic"), ("RBF\nthin-plate", "base_grid8x8_rbf_tps")]
    st = [dev_stats(r) for _, r in runs]
    fig, ax = plt.subplots(figsize=(8, 4.2))
    x = np.arange(len(runs))
    w = 0.38
    ax.bar(x - w / 2 - 0.01, [s[0] for s in st], w, color=C1, label="PSNR medio (3 g)")
    ax.bar(x + w / 2 + 0.01, [s[1] for s in st], w, color=C2, label="PSNR del peor g")
    for i, s in enumerate(st):
        ax.text(x[i] - w / 2, s[0] + 0.3, "%.1f" % s[0], ha="center", fontsize=9)
    ax.set_xticks(x)
    ax.set_xticklabels([l for l, _ in runs])
    ax.set_ylim(10, 28)
    ax.set_ylabel("PSNR [dB]")
    ax.grid(axis="x", visible=False)
    ax.legend(loc="upper left")
    ax.set_title("Interpolación de scipy sobre una grilla 8×8 (64 puntos)", loc="left", fontsize=12)
    return _save(fig, "baselines.png")


def fig_profiles():
    fig, ax = plt.subplots(1, 2, figsize=(12, 4.3), gridspec_kw={"width_ratios": [1, 1.6]})
    t = load_truth("-2.0")
    ax[0].imshow(t, cmap="viridis", vmin=0, vmax=1)
    ax[0].axvline(20, color="w", lw=1.5, ls="--")
    ax[0].set_title("g = -2.0  (corte en la columna 20)", fontsize=11)
    ax[0].set_xticks([])
    ax[0].set_yticks([])
    ax[0].grid(False)
    for g, c in zip(["-4.0", "-2.0", "-0.5"], [C1, C2, C3]):
        ax[1].plot(load_truth(g)[:, 20], color=c, lw=2, label="g = %s" % g)
    ax[1].set_xlabel("fila (hacia abajo en la imagen)")
    ax[1].set_ylabel("valor del mapa")
    ax[1].annotate("meseta ≈ 1", (15, 0.97), fontsize=10, color=INK2)
    ax[1].annotate("rampa suave\n(ancha si g → 0)", (62, 0.55), fontsize=10, color=INK2)
    ax[1].annotate("salto a 0 exacto\n(acantilado)", (100, 0.35), fontsize=10, color=INK2)
    ax[1].legend(loc="lower left")
    ax[1].set_title("Mismo perfil en todos los g: meseta → rampa → acantilado", loc="left", fontsize=11)
    return _save(fig, "perfiles.png")


def fig_sampling(run="VAL9_loo_mono_adapt", g="-2.0"):
    from autoexp import sampling
    from autoexp.oracle import Oracle
    q = json.load(open(os.path.join(RUNS, run, "g%s" % g, "queries.json")))
    P = np.array(q["points"])
    o = Oracle(g)
    sampling.grid(o, 30, nx=6)
    n_grid = o.n_used
    o = Oracle(g)
    sampling.bisect(o, 64, n1=30, nx=6, target="zero", fill="none")
    n_bis = o.n_used
    t = load_truth(g)
    fig, ax = plt.subplots(1, 2, figsize=(12, 5.2))
    ax[0].imshow(t, cmap="gray", vmin=-0.6, vmax=1.3)
    for (a, b, c, lab, m) in [(0, n_grid, C1, "1) grilla (%d)" % n_grid, "o"),
                              (n_grid, n_bis, C2, "2) bisección del acantilado (%d)" % (n_bis - n_grid), "s"),
                              (n_bis, len(P), C3, "3) validación cruzada (%d)" % (len(P) - n_bis), "D")]:
        ax[0].scatter(P[a:b, 1], P[a:b, 0], s=34, color=c, marker=m, edgecolors="white", linewidths=0.8, label=lab)
    ax[0].set_title("g = %s: dónde se consultan los 64 puntos" % g, fontsize=11)
    ax[0].legend(loc="upper right", fontsize=9, framealpha=0.9, frameon=True)
    ax[0].set_xticks([])
    ax[0].set_yticks([])
    ax[0].grid(False)
    rec = np.load(os.path.join(RUNS, run, "g%s" % g, "restored.npy"))
    ax[1].imshow(rec, cmap="viridis", vmin=0, vmax=1)
    ax[1].set_title("reconstrucción (sin DIP): %.1f dB" % per_g(run)[g], fontsize=11)
    ax[1].set_xticks([])
    ax[1].set_yticks([])
    ax[1].grid(False)
    return _save(fig, "muestreo_%s.png" % g)


def fig_val9():
    new = per_g("VAL9_loo_mono_adapt")
    import csv
    rows = {r["g"]: r for r in csv.DictReader(open(os.path.join(REPO, "results", "comparacion_nuevo_vs_dip_v2", "psnr.csv")))}
    d164 = [float(rows[g]["dip_164"]) for g in GS9]
    d328 = [float(rows[g]["dip_328"]) for g in GS9]
    nw = [new[g] for g in GS9]
    fig, ax = plt.subplots(figsize=(13, 4.8))
    x = np.arange(len(GS9))
    w = 0.26
    ax.bar(x - w - 0.02, nw, w, color=C1, label="Nuevo: 64 puntos, sin DIP (solo consulta puntos)")
    ax.bar(x, d164, w, color=C2, label="DIP previo, ~164 puntos (frontier_mix, usa el mapa denso)")
    ax.bar(x + w + 0.02, d328, w, color=C3, label="DIP previo, ~328 puntos (frontier_mix)")
    for i in range(len(x)):
        ax.text(x[i] - w - 0.02, nw[i] + 0.4, "%.1f" % nw[i], ha="center", fontsize=8.5)
    ax.axhline(38, color=INK, ls=(0, (4, 3)), lw=1)
    ax.set_xticks(x)
    ax.set_xticklabels(["g = %s" % g for g in GS9])
    ax.set_ylim(20, 58)
    ax.set_ylabel("PSNR [dB]")
    ax.grid(axis="x", visible=False)
    ax.legend(loc="upper left", fontsize=9, ncol=1)
    ax.set_title("Validación en los 9 g negativos", loc="left", fontsize=12)
    return _save(fig, "val9_vs_dip.png")


def fig_resumen():
    """Una sola imagen con el estado actual: recorrido + reconstrucciones por g."""
    fig = plt.figure(figsize=(16, 8.2))
    gsp = fig.add_gridspec(2, 4, height_ratios=[2.1, 1], hspace=0.3, wspace=0.08, top=0.92)
    fig_progress(fig.add_subplot(gsp[0, :]))
    run = "VAL9_loo_mono_adapt"
    d = per_g(run)
    for k, g in enumerate(["-4.0", "-2.0", "-1.5", "-0.5"]):
        sub = gsp[1, k].subgridspec(1, 2, wspace=0.03)
        a0, a1 = fig.add_subplot(sub[0]), fig.add_subplot(sub[1])
        a0.imshow(load_truth(g), cmap="viridis", vmin=0, vmax=1)
        a1.imshow(np.load(os.path.join(RUNS, run, "g%s" % g, "restored.npy")), cmap="viridis", vmin=0, vmax=1)
        a0.set_title("g=%s original" % g, fontsize=9.5)
        a1.set_title("64 pts: %.1f dB" % d[g], fontsize=9.5)
        for a in (a0, a1):
            a.set_xticks([])
            a.set_yticks([])
            a.grid(False)
    m9 = np.mean([d[g] for g in GS9])
    fig.suptitle("Diagrama de fases con 64 puntos consultados: %.1f dB de media en los 9 g (sin DIP); "
                 "fusión TPS + DIP: %.1f dB en los g de desarrollo" % (m9, dev_stats("fuse2_hib_mid_B10")[0]),
                 fontsize=13, x=0.5, y=0.99)
    return _save(fig, "resumen.png")


# --------------------------------------------------------- helpers pptx --------
def _title(slide, text, size=26):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = BLUE


def add_title_slide(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(12), Inches(2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = BLUE
    tb2 = s.shapes.add_textbox(Inches(0.7), Inches(4.0), Inches(12), Inches(1.5))
    tb2.text_frame.word_wrap = True
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(19)
    p2.font.color.rgb = GREY


def add_bullets_slide(prs, title, bullets, body_size=17):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, title, 27)
    body = s.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12), Inches(5.6))
    tf = body.text_frame
    tf.word_wrap = True
    for j, b in enumerate(bullets):
        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        lvl, text = (1, b[2:]) if b.startswith("- ") else (0, b)
        para.text = text
        para.level = lvl
        para.font.size = Pt(body_size if lvl == 0 else body_size - 3)
        para.font.color.rgb = GREY
        para.space_after = Pt(5)


def _fit(img, max_w, max_h):
    with Image.open(img) as im:
        pw, ph = im.size
    w = float(max_w)
    h = w * ph / pw
    if h > max_h:
        h = float(max_h)
        w = h * pw / ph
    return w, h


def add_image_slide(prs, title, img, caption=None, width=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, title)
    top, bot = 1.4, (6.7 if caption else 7.25)
    w, h = _fit(img, float(width) if width else 12.6, bot - top)
    s.shapes.add_picture(img, Inches((SLIDE_W - w) / 2), Inches(top + (bot - top - h) / 2),
                         width=Inches(w), height=Inches(h))
    if caption:
        cb = s.shapes.add_textbox(Inches(0.6), Inches(bot + 0.05), Inches(12.1), Inches(0.7))
        cb.text_frame.word_wrap = True
        cp = cb.text_frame.paragraphs[0]
        cp.text = caption
        cp.font.size = Pt(12)
        cp.font.italic = True
        cp.font.color.rgb = GREY
        cp.alignment = PP_ALIGN.CENTER


def add_text_image_slide(prs, title, bullets, img, img_w=7.4, body_size=14):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, title)
    body = s.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12.9 - img_w - 0.8), Inches(5.8))
    tf = body.text_frame
    tf.word_wrap = True
    for j, b in enumerate(bullets):
        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        lvl, text = (1, b[2:]) if b.startswith("- ") else (0, b)
        para.text = text
        para.level = lvl
        para.font.size = Pt(body_size if lvl == 0 else body_size - 2)
        para.font.color.rgb = GREY
        para.space_after = Pt(5)
    w, h = _fit(img, img_w, 5.8)
    s.shapes.add_picture(img, Inches(SLIDE_W - w - 0.4), Inches(1.4 + (5.8 - h) / 2), width=Inches(w), height=Inches(h))


def add_table_slide(prs, title, headers, rows, note=None, col_w=None, font=12, left=0.8):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, title)
    nr, nc = len(rows) + 1, len(headers)
    th = min(0.42 * nr, 5.0)
    tbl = s.shapes.add_table(nr, nc, Inches(left), Inches(1.5), Inches(11.7), Inches(th)).table
    if col_w:
        for c, w in enumerate(col_w):
            tbl.columns[c].width = Inches(w)
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.text_frame.paragraphs[0].font.size = Pt(font + 1)
        cell.text_frame.paragraphs[0].font.bold = True
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(v)
            cell.text_frame.paragraphs[0].font.size = Pt(font)
            cell.text_frame.paragraphs[0].font.bold = (c == 0)
    if note:
        nb = s.shapes.add_textbox(Inches(left), Inches(1.6 + th), Inches(11.7), Inches(1.4))
        tf = nb.text_frame
        tf.word_wrap = True
        for j, line in enumerate(note.split("\n")):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            para.text = line
            para.font.size = Pt(12)
            para.font.italic = True
            para.font.color.rgb = GREY


# ----------------------------------------------------------------- build -------
def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    f_prog = fig_progress()
    f_base = fig_baselines()
    f_prof = fig_profiles()
    f_samp = fig_sampling()
    f_val9 = fig_val9()
    f_res = fig_resumen()
    v9 = per_g("VAL9_loo_mono_adapt")
    m9 = np.mean([v9[g] for g in GS9])
    n38 = sum(v9[g] >= 38 for g in GS9)
    fu = dev_stats("fuse2_hib_mid_B10")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)

    add_title_slide(prs, "Reconstruir el diagrama de fases con 64 puntos",
                    "Loop de experimentación automática: muestreo activo + interpolación + DIP  ·  "
                    "Santiago Afonso  ·  (en curso: se actualiza con cada resultado nuevo)")

    add_image_slide(prs, "Resumen: dónde estamos", f_res,
                    "Arriba: cada paso, con el mismo presupuesto de 64 puntos. Abajo: original vs. "
                    "reconstrucción de 64 puntos, sin DIP, en 4 de los 9 g.")

    add_bullets_slide(prs, "El problema y el objetivo", [
        "Cada punto del diagrama con KMC cuesta horas: la pregunta de la tesis es con cuántos "
        "puntos se puede reconstruir el mapa completo de forma confiable.",
        "Objetivo fijado: ≤ 64 puntos (0.4 % de una imagen de 128×128) y > 38 dB de PSNR.",
        "Hasta la entrega anterior, DIP con ~82 puntos daba 14 a 36 dB según el g, y hacían falta "
        "~164 a 328 puntos (ubicados mirando el mapa denso) para pasar los 40 dB.",
        "Regla clave: ningún método puede mirar el mapa. Solo puede consultar puntos, igual "
        "que se haría con KMC.",
    ])

    add_bullets_slide(prs, "Cómo se armó el loop automático", [
        "Oráculo (fijo): la única forma de ver valores del mapa. Cuenta las consultas, corta en "
        "64 y las registra.",
        "Evaluador (fijo): PSNR / SSIM / PSNR en la banda de frontera; verifica el presupuesto y "
        "que los valores registrados sean los reales. Deja una figura por corrida.",
        "Desarrollo con 3 g (-4, -2, -0.5); todo lo que gana se valida en los 9 g negativos.",
        "Cada intento cambia una sola cosa, se mide y se anota en la bitácora "
        "(autoexp/experimentos.md). Si empeora, se revierte.",
        "Interpolación y muestreo corren en local (segundos); DIP corre en Mendieta (A30, ~20 min "
        "por trial). Optuna busca hiperparámetros de DIP en paralelo.",
    ])

    add_text_image_slide(prs, "Punto de partida: interpolación de scipy (punto 6 de la reunión)", [
        "Grilla 8×8 = 64 puntos, cuatro interpoladores de scipy.",
        "El mejor es la RBF thin-plate spline (TPS): 24.4 dB de media, 21.2 en el peor g.",
        "Nota: el interp2d del link ya no existe en scipy moderno; se usaron griddata y "
        "RBFInterpolator, que son sus reemplazos.",
        "Con los mismos 64 puntos, DIP con la receta de la tesis da 25.0: prácticamente igual.",
    ], f_base, img_w=7.6)

    add_image_slide(prs, "La observación que cambió todo", f_prof,
                    "En todos los g, el perfil vertical es meseta → rampa suave → salto a 0 EXACTO. "
                    "Debajo del acantilado el mapa vale 0; arriba es suave.")

    add_text_image_slide(prs, "Muestreo: dónde gastar los 64 puntos", [
        "1) Grilla de 30 puntos (6 columnas): idea global y en qué columnas hay salto a 0.",
        "2) Bisección: en cada columna con salto, búsqueda binaria vertical ('¿vale 0 o no?'). "
        "Ubica el acantilado a 1 px con ~5 consultas por columna.",
        "3) Relleno por validación cruzada: cada punto se predice con los demás y los puntos "
        "nuevos van donde el modelo más se equivoca (tandas de 4).",
        "Con KMC: son ~10 rondas de corridas en paralelo, no 64 corridas en serie.",
    ], f_samp, img_w=7.0, body_size=13)

    add_bullets_slide(prs, "Reconstrucción: interpolar respetando la forma del mapa", [
        "Acantilado: recta ajustada a los pares vecinos 0 / no-0 (robusta a outliers). Debajo: 0.",
        "Arriba: TPS solo con los puntos de arriba (no interpola a través del salto).",
        "- Rampa: cerca del borde, TPS en coordenadas alineadas al borde (el perfil se traslada "
        "paralelo al frente). Ancho de la franja adaptativo según lo abrupto del frente, medido "
        "con el valor justo arriba del acantilado (g=-4: 0.81; g suaves: ~0.10).",
        "- Zona que se desvanece (x ≈ 85): distancia vertical comprimida (depende casi solo de x).",
        "Monotonía: el mapa real nunca sube hacia abajo ni hacia la derecha (0 violaciones en los "
        "9 g; SoC_max baja al crecer ℓ y Ξ).",
        "- Se recorta a las cotas exactas que dan los puntos consultados y se proyecta sobre "
        "mapas monótonos (Dykstra + regresión isotónica). Borra las ondulaciones de la TPS.",
    ], body_size=16)

    add_image_slide(prs, "El recorrido, paso a paso", f_prog,
                    "Media y peor g en los 3 g de desarrollo. De 24.4 dB (grilla + TPS) a %.1f dB "
                    "(fusión con DIP)." % fu[0])

    add_image_slide(prs, "Validación en los 9 g: %.1f dB de media con 64 puntos, sin DIP" % m9, f_val9,
                    "%d de 9 g ≥ 38 dB. Le gana a DIP con ~164 puntos en los g duros (-4, -3.5, -3); "
                    "DIP con 164 sigue siendo mejor en los g suaves." % n38)

    add_image_slide(prs, "Comparación visual con el mejor DIP previo",
                    os.path.join(REPO, "results", "comparacion_nuevo_vs_dip_v2", "comparacion_rec.png"),
                    width=6.2)
    add_image_slide(prs, "Mapas de error (misma grilla)",
                    os.path.join(REPO, "results", "comparacion_nuevo_vs_dip_v2", "comparacion_err.png"),
                    width=6.2)

    rob = [(o, score(r)) for o, r in [("0.20", "ROBUST9_off0.2"), ("0.35", "VAL9_loo_n30_off0.35"),
                                      ("0.50", "VAL9_loo_n30_off0.5"), ("0.65", "VAL9_loo_n30_off0.65"),
                                      ("0.80", "ROBUST9_off0.8")]]
    add_table_slide(prs, "Robustez: ¿depende de dónde cae la grilla?",
                    ["offset de la grilla", "media 9 g [dB]", "peor g [dB]"],
                    [[o, "%.1f" % s["mean_psnr"], "%.1f" % s["min_psnr"]] for o, s in rob if s],
                    note="Todo se afinó con offset 0.5, así que ese valor está algo favorecido. Lo honesto "
                         "es reportar la banda (~37.5 ± 1 dB con esta versión del muestreo), no el mejor "
                         "caso.\n(Filas 0.20 y 0.80: versión anterior del relleno.)",
                    col_w=[3.9, 3.9, 3.9])

    add_table_slide(prs, "DIP con 64 puntos: hiperparámetros del paper vs. los de la tesis",
                    ["configuración (grilla 8×8, 64 pts)", "media [dB]", "peor [dB]"],
                    [["interpolación TPS (referencia)", "24.4", "21.2"],
                     ["DIP, receta de la tesis (LR 1e-3, reg 0.08, skip 4, bilinear)", "25.0", "21.9"],
                     ["DIP, paper 'kate' (noise 32, LR 0.01, skip 128, nearest)", "18.1", "13.7"],
                     ["DIP, paper 'vase' (meshgrid, LR 0.01, reg 0.03, skip 0, nearest)", "28.5", "25.3"],
                     ["DIP, 'vase' + muestreo uniforme", "22.2", "19.4"]],
                    note="Con los hiperparámetros del paper para agujeros grandes, DIP le gana por +4 dB "
                         "a la interpolación con los mismos puntos; la receta de la tesis (afinada para "
                         "~328 puntos) no. Estudio 'profes_grid_uniform' de Optuna en Mendieta.",
                    col_w=[7.3, 2.2, 2.2])

    add_table_slide(prs, "Híbrido: TPS pegada al borde + DIP en la zona suave",
                    ["reconstrucción (mismos 64 puntos, 3 g)", "media [dB]", "peor [dB]", "-4 / -2 / -0.5"],
                    [["TPS sola (con monotonía)", "40.4", "36.8", "45.7 / 36.8 / 38.5"],
                     ["DIP (tesis) + ceros + franja TPS", "39.3", "38.0", "41.7 / 38.0 / 38.1"],
                     ["DIP ('vase') + ceros + franja TPS", "35.4", "30.8", "30.8 / 36.5 / 38.9"],
                     ["Fusión TPS + DIP tesis (B = 20 px)", "41.8", "39.5", ""],
                     ["Fusión TPS + DIP intermedio (B = 10 px)",
                      "%.1f" % fu[0], "%.1f" % fu[1], " / ".join("%.1f" % per_g("fuse2_hib_mid_B10")[g] for g in DEV)]],
                    note="Fusión: final = monotonía( w·TPS + (1−w)·DIP ), w = exp(−(d/B)²), d = distancia "
                         "al acantilado. DIP va mejor en la meseta y en la zona suave; la TPS, en el frente "
                         "abrupto. Juntas, los 3 g quedan ≥ 39.5 dB. Validación en 9 g: en curso.",
                    col_w=[5.0, 1.8, 1.8, 3.1])

    add_image_slide(prs, "Punto 5 de la reunión: Rosenbrock de vuelta a 3D",
                    os.path.join(REPO, "results", "rosenbrock_3d", "rosenbrock_3d.png"),
                    "Con 1 % DIP recupera la forma pero la cresta angosta queda ondulada; con 5 % "
                    "queda fiel. Con 64 puntos la TPS da una superficie reconocible pero con bultos.")

    add_bullets_slide(prs, "Pedidos de la reunión: estado", [
        "1–3) Optuna + grid/uniform con hiperparámetros + la grilla como hiperparámetro: estudio "
        "en Mendieta. La config 'vase' del paper es la mejor para DIP con 64 puntos. En el muestreo "
        "nuevo, la grilla (6 columnas × 5 filas) resultó ser el parámetro más sensible.",
        "4) Ventana: la bisección ya concentra ~25 de los 64 puntos en una franja angosta "
        "alrededor del frente (ventana adaptativa).",
        "5) Rosenbrock en 3D: hecho (diapositiva anterior).",
        "6) Interpolación de scipy vs. DIP: con los mismos puntos empatan con la receta de la "
        "tesis (24.4 vs. 25.0); DIP con la config del paper gana (28.5).",
    ], body_size=16)

    add_bullets_slide(prs, "Lo que se probó y NO funcionó (para no repetirlo)", [
        "Frente como parábola por puntos medios de pares con salto: inestable, dejaba una costura.",
        "Bisección al nivel medio: en los g suaves caía dentro de la rampa, no en el acantilado.",
        "Sondas del perfil de la rampa (3-7-14 px sobre el borde): le quitan presupuesto al relleno.",
        "Proceso gaussiano en lugar de la TPS: con tan pocos puntos, las escalas aprendidas "
        "quedan mal (g=-2: 22 dB vs. 36).",
        "Relleno por ancho de las cotas de monotonía: mejor peor g, pero −2.7 dB en g=-4.",
        "Grilla de 7-8 columnas: no falla el detector, se acaba el presupuesto en la bisección.",
        "DIP 'kate' del paper con 64 puntos (18 dB): pensado para imágenes casi completas.",
    ], body_size=15)

    add_bullets_slide(prs, "Para KMC: recomendación y lo que falta validar", [
        "Muestreo por tandas: grilla (paralelo) → ~5 rondas de bisección (una corrida por columna, "
        "en paralelo) → ~4 rondas de relleno de a 4. Unas 10 rondas de KMC en total.",
        "Elegir las columnas según el presupuesto, dejando ~10 puntos para el relleno.",
        "Reconstrucción: fusión TPS (cerca del acantilado) + DIP (zona suave) + monotonía.",
        "Pendiente, y lo más importante:",
        "- ¿Existe el acantilado a 0 en KMC, o el borde es difuso / ruidoso?",
        "- Ruido: KMC es estocástico. La TPS pasa exacto por cada punto (copia el ruido); DIP es un "
        "prior pensado para imágenes ruidosas. Hay que medirlo con ruido simulado del tamaño del de KMC.",
        "- Validar la fusión en los 9 g (DIP corriendo en Mendieta).",
    ], body_size=16)

    out = os.path.join(OUT_DIR, PRES_NAME + ".pptx")
    prs.save(out)
    print("->", out, "(%d diapositivas)" % len(prs.slides))
    print("->", f_res)


if __name__ == "__main__":
    main()
