#!/usr/bin/env python
# coding: utf-8
"""Retoques puntuales sobre presentacion1_02_09.pptx, sin tocar el resto:
  1. reescribe la slide 3 (la del codigo) -> "Donde se calcula cada metrica"
  2. agrega una slide simple de contexto como slide 1

Uso:  ./venv/bin/python presentaciones/presentacion1_02_09/fix_slide3_y_contexto.py
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, "presentacion1_02_09.pptx")

BLUE = RGBColor(0x1F, 0x3B, 0x73)
GREY = RGBColor(0x40, 0x40, 0x40)
CODEBG = RGBColor(0xF3, 0xF3, 0xF3)

CODE = """\
from skimage.metrics import peak_signal_noise_ratio as compare_psnr
from skimage.metrics import structural_similarity  as compare_ssim

# cada 100 iteraciones (mientras entrena)
psrn        = compare_psnr(img_np, out_np)                    # PSNR
psrn_masked = compare_psnr(img_masked, out_np * img_mask_np)  # PSNR (solo lo observado)
ssim_full   = compare_ssim(img_np[0], out_clip[0], data_range=1.0)   # SSIM

# una vez, al terminar
abs_err    = np.abs(img_np - out_np)      # mapa de error por pixel
final_psnr = compare_psnr(img_np, out_np)
final_ssim = compare_ssim(img_np[0], out_np[0], data_range=1.0)
final_mae  = float(abs_err.mean())        # MAE = promedio del mapa"""

BULLETS = [
    "PSNR  ->  funcion de scikit-image (skimage.metrics)",
    "SSIM  ->  funcion de scikit-image (skimage.metrics)",
    "MAE   ->  a mano con numpy: np.abs(orig - rec).mean()",
]

CTX_TITLE = "Que me pidieron esta semana"
CTX_BULLETS = [
    "Comparar la reconstruccion DIP del diagrama de fases con distintos "
    "porcentajes de mascara.",
    "Mirar algunas metricas de error para ponerle numero a esa reconstruccion.",
    "Imagen de partida: el mapa del modelo del continuo (galpynostatic), 128x128.",
    "Mascaras probadas: 50, 70, 80, 90, 95, 98 y 99 % de puntos ocultos.",
    "Metricas: PSNR, SSIM y MAE.",
]


def clear(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def add_title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = BLUE


def rebuild_slide3(slide):
    clear(slide)
    add_title(slide, "Donde se calcula cada metrica (restorationGRIS.py)")

    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(12), Inches(4.2))
    box.fill.solid()
    box.fill.fore_color.rgb = CODEBG
    tf = box.text_frame
    tf.word_wrap = True
    for j, line in enumerate(CODE.split("\n")):
        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        para.text = line
        para.font.name = "Consolas"
        para.font.size = Pt(12)
        para.font.color.rgb = GREY

    bb = slide.shapes.add_textbox(Inches(0.7), Inches(5.8), Inches(12), Inches(1.4))
    tfb = bb.text_frame
    tfb.word_wrap = True
    for j, b in enumerate(BULLETS):
        para = tfb.paragraphs[0] if j == 0 else tfb.add_paragraph()
        para.text = "- " + b
        para.font.size = Pt(14)
        para.font.color.rgb = GREY
        para.space_after = Pt(3)


def build_context(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, CTX_TITLE)
    body = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(12), Inches(4.5))
    tf = body.text_frame
    tf.word_wrap = True
    for j, b in enumerate(CTX_BULLETS):
        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        para.text = "- " + b
        para.font.size = Pt(18)
        para.font.color.rgb = GREY
        para.space_after = Pt(10)
    return s


def move_first(prs, from_idx):
    lst = prs.slides._sldIdLst
    el = list(lst)[from_idx]
    lst.remove(el)
    lst.insert(0, el)


def main():
    prs = Presentation(PPTX)

    # localizar la slide del codigo (contiene 'compare_psnr')
    code_idx = None
    for i, sl in enumerate(prs.slides):
        txt = " ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
        if "compare_psnr" in txt:
            code_idx = i
            break
    if code_idx is None:
        print("No encontre la slide del codigo, no toco nada.")
        return
    rebuild_slide3(prs.slides[code_idx])
    print("Slide del codigo (posicion %d) reescrita." % (code_idx + 1))

    # contexto como slide 1 (si no existe ya)
    already = any(
        sh.has_text_frame and sh.text_frame.text.strip().startswith(CTX_TITLE)
        for sl in prs.slides for sh in sl.shapes
    )
    if already:
        print("La slide de contexto ya existe, no se agrega.")
    else:
        build_context(prs)
        move_first(prs, len(prs.slides._sldIdLst) - 1)
        print("Slide de contexto agregada como slide 1.")

    prs.save(PPTX)
    print("Guardado:", PPTX, "(%d diapositivas)" % len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
