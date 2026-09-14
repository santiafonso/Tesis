#!/usr/bin/env python
# coding: utf-8
"""Agrega UNA diapositiva a presentacion1_02_09.pptx (la que ya venis editando a
mano) sin tocar ninguna de las diapositivas existentes: "Donde se usa cada
metrica en el codigo".  La inserta justo despues de la slide 3 (la del codigo).

Uso:  ./venv/bin/python presentaciones/presentacion1_02_09/add_slide_uso_metricas.py
Es idempotente: si la slide ya existe, no la duplica.
"""
import copy
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, "presentacion1_02_09.pptx")

BLUE = RGBColor(0x1F, 0x3B, 0x73)
GREY = RGBColor(0x40, 0x40, 0x40)
RED = RGBColor(0xB0, 0x2A, 0x2A)
CODEBG = RGBColor(0xF3, 0xF3, 0xF3)
HDRBG = RGBColor(0x1F, 0x3B, 0x73)
HL = RGBColor(0xFC, 0xEF, 0xC7)

TITLE = "Donde se usa cada metrica en el codigo"

ROWS = [
    # (metrica, como se calcula, para que se usa, resaltar?)
    ("PSNR  (imagen completa)\npsrn",
     "skimage.metrics\ncompare_psnr(img_np, out_np)",
     "Solo reporte: consola  ·  metrics.csv (col. psnr_full)  ·  "
     "curva psnr_curve.png  ·  titulo de comparison_annotated.png",
     False),
    ("PSNR  (solo observados)\npsrn_masked",
     "skimage.metrics\ncompare_psnr(img_masked,\n            out_np * img_mask_np)",
     "UNICA que ACTUA sobre el entrenamiento: si cae mas de 5 dB respecto "
     "a la ultima medicion -> vuelve al checkpoint anterior (backtracking).",
     True),
    ("SSIM  (imagen completa)\nssim_full",
     "skimage.metrics\ncompare_ssim(img_np[0], out_np[0],\n             data_range=1.0)",
     "Solo reporte: consola  ·  metrics.csv (col. ssim_full)  ·  "
     "curva psnr_curve.png  ·  titulo de comparison_annotated.png",
     False),
    ("MAE\nfinal_mae",
     "numpy (no hay libreria)\nabs_err = np.abs(img_np - out_np)\n"
     "final_mae = abs_err.mean()",
     "Solo reporte/visual: mapa de error |error| en final_comparison.png y "
     "comparison_annotated.png  ·  titulo de ese panel  ·  metrics.csv (fila mae).",
     False),
]

FOOT = ("En resumen: psrn_masked es la unica metrica que cambia lo que hace el "
        "algoritmo (backtracking). PSNR, SSIM y MAE de la imagen completa son solo "
        "para medir y mostrar el resultado.")


def build_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])

    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = TITLE
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = BLUE

    rows, cols = len(ROWS) + 1, 3
    tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.35),
                             Inches(12.35), Inches(4.7)).table
    tbl.columns[0].width = Inches(2.7)
    tbl.columns[1].width = Inches(3.9)
    tbl.columns[2].width = Inches(5.75)

    for c, h in enumerate(["Metrica", "Como se calcula", "Para que se usa"]):
        cell = tbl.cell(0, c)
        cell.text = h
        pr = cell.text_frame.paragraphs[0]
        pr.font.size = Pt(13)
        pr.font.bold = True
        pr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.fill.solid()
        cell.fill.fore_color.rgb = HDRBG

    for r, (met, how, use, hl) in enumerate(ROWS, start=1):
        for c, val in enumerate((met, how, use)):
            cell = tbl.cell(r, c)
            cell.text = val
            for k, para in enumerate(cell.text_frame.paragraphs):
                para.font.size = Pt(10.5 if c else 11)
                para.font.color.rgb = RED if hl else GREY
                if c == 0:
                    para.font.bold = (k == 0)
                if c == 1:
                    para.font.name = "Consolas"
            cell.fill.solid()
            cell.fill.fore_color.rgb = HL if hl else RGBColor(0xFF, 0xFF, 0xFF)

    fb = s.shapes.add_textbox(Inches(0.6), Inches(6.25), Inches(12.1), Inches(1.0))
    fp = fb.text_frame
    fp.word_wrap = True
    par = fp.paragraphs[0]
    par.text = FOOT
    par.font.size = Pt(12)
    par.font.italic = True
    par.font.color.rgb = GREY
    par.alignment = PP_ALIGN.CENTER
    return s


def move_after(prs, from_idx, after_idx):
    """Mueve la slide en from_idx para que quede justo despues de after_idx."""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    el = ids[from_idx]
    sldIdLst.remove(el)
    ids = list(sldIdLst)
    sldIdLst.insert(after_idx + 1, el)


def main():
    prs = Presentation(PPTX)
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(TITLE):
                print("La diapositiva ya existe, no se agrega de nuevo.")
                return
    build_slide(prs)                       # queda al final
    move_after(prs, len(prs.slides._sldIdLst) - 1, 2)   # -> posicion 4 (tras la del codigo)
    prs.save(PPTX)
    print("Agregada 'Donde se usa cada metrica' como diapositiva 4 en:", PPTX,
          "(%d diapositivas)" % len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
