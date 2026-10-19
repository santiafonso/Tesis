#!/usr/bin/env python
# coding: utf-8
"""Genera presentacion7_10_10.pptx: avance de la reconstruccion del diagrama de fases con KMC en los
13 g (-8 a 4), para ir mostrando a los profes. Mismo estilo que la 5 y la 6.

Los g se muestran con KMC; el continuo aparece solo como referencia (y en la resta KMC - continuo).

Uso (desde ~/Tesis-autoexp, con el estado de Mendieta copiado en autoexp/kmc/):
    ./venv/bin/python presentaciones/presentacion7_10_10/figuras.py
    ./venv/bin/python presentaciones/presentacion7_10_10/make_presentacion7_10_10.py
"""
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(REPO, "presentaciones", "explicacion_autoexp"))
sys.path.insert(0, os.path.join(REPO, "presentaciones", "presentacion5_05_10"))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import make_presentacion5_05_10 as P5  # noqa: E402  (estilo de reunion: titulo azul, texto gris)

M = P5.M
FIG = os.path.join(HERE, "fig")
OUT = os.path.join(HERE, "presentacion7_10_10.pptx")


def gtxt(g):
    return "g = %+d" % g if g else "g = 0"


def s_portada(prs):
    s = M.new_slide(prs)
    P5._text(s, 0.7, 2.2, 12.0, 2.0, ["Reconstrucción del diagrama de fases con KMC: avance"], size=36,
             color=P5.TBLUE)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    P5._text(s, 0.7, 4.0, 12.0, 1.5, ["13 valores de g (−8 a 4), ~96 puntos de KMC por g, control con puntos no "
                                      "usados  ·  10/10  ·  Santiago Afonso"], size=19, color=P5.TGREY)


def s_estado(prs, R):
    term = [gtxt(r["g"]) for r in R if r["etapa"] == "terminado"]
    s = M.new_slide(prs)
    M.title(s, "Dónde está cada g", sub="puntos de KMC por g · el loop decide solo dónde medir")
    M.image(s, os.path.join(FIG, "estado_g.png"), 0.4, 1.55, w=8.6, h=5.7)
    M.card(s, 9.25, 1.7, 3.5, 5.4, "Resumen", [
        "Terminados (mapa final + control): %s." % ", ".join(term),
        "g = −5: mapa listo, control corriendo.",
        "Los otros 7 siguen eligiendo puntos: tandas de 3-4 en el borde.",
        "Cotas: corridas cortadas por tiempo (log ℓ bajo); se usan como \"vale al menos esto\"."], size=13)
    M.notes(s, """
Cada g tiene su loop: grilla de 30 puntos, bisección sobre el borde y relleno donde más falta. Termina cuando
el planificador no pide más (no exactamente en 96). Las cotas son corridas que no terminan en el tiempo del
cluster (log ℓ ≈ −3.5 tarda días): no son el valor, son un mínimo.
""")


def s_finales(prs):
    s = M.new_slide(prs)
    M.title(s, "Mapas finales: g = −8 a −5",
            sub="arriba KMC (TPS + DIP) con sus puntos · medio el continuo, solo referencia · abajo la diferencia")
    M.image(s, os.path.join(FIG, "finales_neg_5a8.png"), 0.4, 1.5, w=12.5, h=5.85)
    M.notes(s, """
Blancos: puntos medidos; círculos rojos: cotas. Lo que muestra la resta (abajo) es la diferencia entre el KMC
y el modelo del continuo, no error de la reconstrucción (el error se mide con el control).
Dos diferencias sistemáticas: (1) el acantilado del KMC está corrido hacia Ξ más bajo (banda roja);
(2) la caída de SoC en log ℓ ≈ −1 a 0 pasa antes (banda azul). Y la franja de log Ξ ≈ 2 con SoC ≈ 0 en g = −8.
En g = −5 los huecos azules arriba a la izquierda son corridas que dan SoC ≈ 0 (la anomalía de Ξ alto).
""")
    s = M.new_slide(prs)
    M.title(s, "Mapas finales: g = −2 y g = +4", sub="mismo formato: KMC · continuo (referencia) · diferencia")
    M.image(s, os.path.join(FIG, "finales_resto.png"), 0.4, 1.45, w=6.2, h=5.9)
    M.card(s, 7.0, 1.7, 5.75, 5.4, "Qué se ve", [
        "g = −2: muy parecido al continuo; el acantilado del KMC un poco más abajo en Ξ (banda roja fina).",
        "g = +4: la meseta del KMC queda en ~0.85, no en 1, y la transición es más ancha (todo azul).",
        "Las cotas (rojos) se concentran en log ℓ < −2.5: ahí el KMC no termina en 48 h.",
        "En g = +4 hay un escalón horizontal en log Ξ ≈ −1.6 a la izquierda: lo arma una cota; a revisar."],
           size=13)
    M.notes(s, """
g = −2 terminó el 10/10 a la mañana; g = +4 a la madrugada. Los dos con 32 puntos de control.
""")


def s_control(prs, R):
    s = M.new_slide(prs)
    M.title(s, "¿Qué tan bien reconstruye? Control con 32 puntos que no se usaron",
            sub="por g: 24 sorteados donde el mapa más cambia o más difiere del continuo + 8 al azar")
    M.image(s, os.path.join(FIG, "control_rms.png"), 0.4, 1.55, w=7.6, h=5.6)
    rr = [r for r in R if r["rms"] is not None]
    M.card(s, 8.3, 1.7, 4.45, 5.4, "Lectura", [
        "Al azar, el error está en el nivel del ruido del KMC (0.002-0.004), salvo g = −6.",
        "En la zona crítica 0.012-0.017, salvo g = −2 (0.035): ahí el error está en el pie del acantilado.",
        "g = −6: un punto al azar con 0.09 de error, donde DIP empeoró (TPS sola: %.3f)." %
        next(r["rms_tps"] for r in rr if r["g"] == -6),
        "La zona crítica es la prueba exigente a propósito."], size=13)
    M.notes(s, """
El control se elige después de reconstruir y se corre con KMC; no entra en la reconstrucción. 3/4 de los
puntos se sortean con peso en |reconstrucción − continuo| + |gradiente|: el frente y lo que más se aparta
del continuo. El RMS de esos es pesimista por diseño; el de los al azar dice cómo está el resto del mapa.
""")
    s = M.new_slide(prs)
    M.title(s, "Ejemplo: control de g = −2", sub="izquierda: mapa final y dónde cayeron los 32 · derecha: reconstruido vs. KMC")
    M.image(s, os.path.join(FIG, "control_g-2.png"), 0.4, 1.6, w=12.5, h=5.6)
    M.notes(s, """
Los críticos (rojo) caen sobre el acantilado y en la franja de log ℓ ≈ 0, los al azar (naranja) casi todos
en la zona de SoC 0. Los dos puntos más errados: SoC 0.08-0.10 en el KMC, ~0.19 en el mapa: pie del acantilado.
""")


def s_en_curso(prs):
    s = M.new_slide(prs)
    M.title(s, "Los g que siguen eligiendo puntos",
            sub="reconstrucción con lo que hay hoy (blancos: medidos · rojos: cotas)")
    M.image(s, os.path.join(FIG, "en_curso.png"), 0.4, 1.5, w=12.5, h=5.85)
    M.notes(s, """
g = −4 y g = +3 son los que primero se cerraron con 64 puntos; ahora van a 96. Los huecos azules arriba a la
izquierda en g = −1 y g = +3 son corridas de KMC que dan SoC ≈ 0 en la zona de Ξ alto (la anomalía); con el
relleno nuevo no se mandan más puntos ahí.
""")


def s_grid(prs):
    s = M.new_slide(prs)
    M.title(s, "Grilla 64×64 de g = −4 con KMC, completa",
            sub="todos los puntos, sin planificador ni reconstrucción · izquierda KMC · centro continuo · derecha la diferencia")
    M.image(s, os.path.join(FIG, "grid64.png"), 0.4, 1.6, w=12.5, h=4.0)
    M.text(s, 0.6, 5.75, 12.1, 1.5, [
        [("4096 puntos: 3597 terminados y 499 cotas ", {"bold": True}),
         ("(log ℓ < −2.9: no terminan en 24 h; triángulos blancos). ", {}),
         ("Confirma lo que ya veían las reconstrucciones: acantilado corrido a Ξ bajo, SoC ≈ 0 en toda la esquina "
          "de Ξ alto y ℓ bajo, y unos pocos píxeles sueltos en 0 dentro de la meseta.", {})]], size=14)
    M.notes(s, """
Pedido de la reunión del 5/10. Las corridas que quedan terminan esta noche (01:00): la mayoría va a salir como
cota. Con checkpoint (filmina siguiente) se podrían terminar las columnas de log ℓ < −3.
Esta imagen sirve de control denso para la reconstrucción de 96 puntos de g = −4 cuando termine.
""")


def s_techo(prs):
    s = M.new_slide(prs)
    M.title(s, "Con g > 0 el KMC no llega a SoC 1", sub="máximo medido en cada g vs. el continuo")
    M.image(s, os.path.join(FIG, "techo_g.png"), 0.4, 1.55, w=6.9, h=5.0)
    M.card(s, 7.6, 1.7, 5.15, 5.4, "¿Es real o es por las cotas?", [
        "El techo baja de forma ordenada: 0.97, 0.95, 0.91, 0.86 para g = 1 a 4.",
        "Pero en g > 0 lo más alto del continuo está en log ℓ < −2.5, donde el KMC solo dio cotas.",
        "La reconstrucción no pasa del máximo medido: si ahí el KMC llega más alto, el mapa no lo puede mostrar.",
        "Para saberlo: correr esos puntos hasta el final (filmina siguiente)."], size=13)
    M.notes(s, """
Hipótesis física (sin verificar): con interacción repulsiva el potencial sube fuerte cerca de la red llena y
toca el corte antes. Las corridas sin tope de la filmina siguiente lo deciden.
""")


def s_checkpoint(prs):
    s = M.new_slide(prs)
    M.title(s, "KMC sin límite de tiempo: checkpoint", sub="para terminar las corridas de log ℓ bajo")
    M.text(s, 0.6, 1.6, 6.2, 5.5, [
        "El cluster corta todo a las 48 h; log ℓ = −3.5 tarda ~2.5 días y log ℓ = −4, ~6.",
        "Ahora el KMC guarda su estado completo (red, generador aleatorio, tasas) cada 10 min y al ser cortado.",
        "Al relanzarlo sigue donde quedó: la corrida cortada dos veces da exactamente el mismo archivo que la "
        "corrida sin cortar (bit a bit).",
        "Las corridas pasan solas de un nodo reservado al siguiente."], size=15, bullets=True, space=10)
    P5.tabla(s, 7.1, 1.7, 5.7, ["log Ξ", "log ℓ", "hoy en el mapa", "qué es"], [
        ["0.20", "−3.53", "0.85", "cota"],
        ["−0.08", "−3.24", "0.85", "cota"],
        ["0.82", "−3.01", "0.86", "cota"],
        ["0.30", "−3.01", "0.83", "cota (0.53)"],
        ["0.68", "−2.54", "0.85", "cota"],
        ["0.82", "−2.02", "0.86", "medido (referencia)"]], col_w=[1.0, 1.0, 1.6, 2.1], font=12)
    M.text(s, 7.1, 4.9, 5.7, 1.5, ["g = +4, 2 réplicas cada uno, arrancaron el 10/10 a las 18:00. Si terminan "
                                   "arriba de 0.86, el techo del mapa está mal puesto."], size=13)
    M.notes(s, """
Se probó local: misma semilla, una corrida entera y otra cortada dos veces y retomada: los datos de salida
tienen el mismo md5 que los del binario original. Sin activar el checkpoint, el KMC es el de siempre.
El último punto ya está medido (0.859): se repite para ver la dispersión con el mismo procedimiento.
""")


def s_sigue(prs):
    s = M.new_slide(prs)
    M.title(s, "Cómo sigue y qué consultar")
    M.text(s, 0.6, 1.6, 12.1, 5.6, [
        [("Cerrar los 7 g que faltan: ", {"bold": True}), ("mapa final (TPS + DIP) y control de 32 puntos en cada uno.", {})],
        [("g = +4 sin tope: ", {"bold": True}), ("si el techo sube, rehacer los g > 0 con esas corridas.", {})],
        [("Grilla 64×64: ", {"bold": True}), ("error de la reconstrucción de 96 puntos de g = −4 contra KMC completo.", {})],
        [("Para consultar: ", {"bold": True}),
         ("(1) la esquina de Ξ alto y ℓ bajo da SoC ≈ 0 en el KMC (continuo ~1), con réplicas que a veces dan 0.99 "
          "y a veces 0; (2) el acantilado del KMC está más abajo en Ξ que en el continuo en todos los g negativos; "
          "(3) el techo con g > 0.", {})]], size=17, bullets=True, space=12)
    M.notes(s, "Próximos pasos y las tres preguntas para los profes.")


def main():
    R = json.load(open(os.path.join(FIG, "resumen.json")))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(M.SW), Inches(M.SH)
    s_portada(prs)
    s_estado(prs, R)
    s_finales(prs)
    s_control(prs, R)
    s_en_curso(prs)
    s_grid(prs)
    s_techo(prs)
    s_checkpoint(prs)
    s_sigue(prs)
    prs.save(OUT)
    print("->", OUT, "(%d filminas)" % len(prs.slides))


if __name__ == "__main__":
    main()
